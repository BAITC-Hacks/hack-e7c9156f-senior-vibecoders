from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from app.services.documents import document_lines, get_clause


class DocumentParsingTest(TestCase):
    def setUp(self) -> None:
        self.temp = TemporaryDirectory()
        self.folder = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_docx_paragraphs_tables_and_clause_boundaries(self) -> None:
        path = self.folder / "source.docx"
        document = Document()
        for text in ("Введение", "3.4. Основной пункт", "Продолжение пункта", "3.5. Следующий пункт"):
            document.add_paragraph(text)
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text = "Отдел"
        row.cells[1].text = "Задача"
        document.save(path)

        self.assertIn("Отдел | Задача", document_lines(path))
        self.assertEqual(get_clause(path, "3.4"), "3.4. Основной пункт\nПродолжение пункта")
        self.assertIsNone(get_clause(path, "9.9"))

    def test_docx_lettered_subclause(self) -> None:
        path = self.folder / "subclauses.docx"
        document = Document()
        document.add_paragraph("3.4. Общий пункт")
        document.add_paragraph("3.4.а. Подпункт А")
        document.add_paragraph("3.4.б. Подпункт Б")
        document.save(path)

        self.assertEqual(get_clause(path, "3.4.а"), "3.4.а. Подпункт А")
        self.assertEqual(get_clause(path, "3.4"), "3.4. Общий пункт")

    def test_xlsx_reads_all_sheets_and_closes_workbook(self) -> None:
        path = self.folder / "source.xlsx"
        workbook = Workbook()
        workbook.active.append(["3.4. Первый пункт", None])
        workbook.active.append(["его продолжение", 7])
        workbook.create_sheet("После").append(["3.5. Второй пункт"])
        workbook.save(path)
        workbook.close()

        self.assertEqual(
            document_lines(path),
            ["3.4. Первый пункт", "его продолжение 7", "3.5. Второй пункт"],
        )
        self.assertEqual(get_clause(path, "3.4"), "3.4. Первый пункт\nего продолжение 7")

    def test_pdf_text_pages_and_scanned_page(self) -> None:
        path = self.folder / "source.pdf"
        pages = [
            SimpleNamespace(extract_text=lambda: "3.4. Первый пункт\n продолжение "),
            SimpleNamespace(extract_text=lambda: None),
            SimpleNamespace(extract_text=lambda: "3.5. Второй пункт"),
        ]
        with patch("app.services.documents.PdfReader", return_value=SimpleNamespace(pages=pages)):
            self.assertEqual(
                document_lines(path),
                ["3.4. Первый пункт", "продолжение", "3.5. Второй пункт"],
            )
            self.assertEqual(get_clause(path, "3.4"), "3.4. Первый пункт\nпродолжение")

    def test_unsupported_document(self) -> None:
        with self.assertRaisesRegex(ValueError, "Неподдерживаемый"):
            document_lines(self.folder / "source.txt")
