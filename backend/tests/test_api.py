import json
from copy import deepcopy
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from app.config import BACKEND_DIR, Settings
from app.main import create_app


def docx_bytes(text: str) -> bytes:
    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


class ApiTest(TestCase):
    def setUp(self) -> None:
        self.temp = TemporaryDirectory()
        config = Settings(_env_file=None, ai_mock=True, storage_dir=Path(self.temp.name))
        self.client = TestClient(create_app(config))

    def tearDown(self) -> None:
        self.client.close()
        self.temp.cleanup()

    def test_demo_end_to_end(self) -> None:
        self.assertEqual(self.client.get("/health").json(), {"status": "ok"})
        response = self.client.post("/api/analyses/demo")
        self.assertEqual(response.status_code, 200)
        analysis_id = response.json()["id"]

        status = self.client.get(f"/api/analyses/{analysis_id}")
        self.assertEqual(status.json()["status"], "done")
        self.assertEqual(status.json()["progress"], 1)

        result_response = self.client.get(f"/api/analyses/{analysis_id}/result")
        self.assertEqual(result_response.status_code, 200)
        result = result_response.json()
        self.assertEqual(result["meta"]["provider"], "mock")
        self.assertTrue(any(change["status"] == "created" for change in result["unit_changes"]))
        self.assertIn("documents", result)
        self.assertIn("alignments", result)
        self.assertIn("flows", result)
        self.assertIn("categories", result)
        self.assertIn("rejected_findings", result)

        evidence_list = [e for item in result["units"] for e in item["evidence"]]
        evidence_list += [e for item in result["unit_changes"] for e in item["evidence"]]
        evidence_list += [e for item in result["findings"] for e in item["evidence"]]
        for evidence in evidence_list:
            clause = self.client.get(
                f"/api/analyses/{analysis_id}/documents/{evidence['doc_id']}/clauses/{evidence['clause_id']}"
            )
            self.assertEqual(clause.status_code, 200)
            self.assertIn(evidence["quote"], clause.json()["text"])

        report = self.client.get(f"/api/analyses/{analysis_id}/report.docx")
        self.assertEqual(report.status_code, 200)
        self.assertTrue(report.content.startswith(b"PK"))
        self.assertIn(f"zaklyuchenie_{analysis_id}.docx", report.headers["content-disposition"])

    def test_review_history_and_report(self) -> None:
        first_id = self.client.post("/api/analyses/demo").json()["id"]
        second_id = self.client.post("/api/analyses/demo").json()["id"]
        history = self.client.get("/api/analyses").json()
        self.assertEqual([row["id"] for row in history], [second_id, first_id])
        self.assertEqual(history[0]["counts"], {"findings": 1, "rejected": 1, "high": 0})
        self.assertEqual([d["side"] for d in history[0]["documents"]], ["before", "after"])
        self.assertIn("created_at", self.client.get(f"/api/analyses/{first_id}").json())
        self.assertEqual(self.client.get(f"/api/analyses/{first_id}/report.docx").status_code, 200)

        accepted = self.client.patch(
            f"/api/analyses/{first_id}/findings/f2",
            json={"status": "accepted", "comment": "Проверено"},
        )
        self.assertEqual(accepted.status_code, 200)
        self.assertIn("Возможная утрата должностей", accepted.json()["conclusion_md"])
        self.assertEqual(accepted.json()["rejected_findings"][0]["review"]["status"], "accepted")
        self.assertIn("reviewed_at", accepted.json()["rejected_findings"][0]["review"])

        rejected = self.client.patch(
            f"/api/analyses/{first_id}/findings/f1", json={"status": "rejected"},
        )
        self.assertEqual(rejected.status_code, 200)
        self.assertNotIn("Созданы ДИТААД", rejected.json()["conclusion_md"])
        self.assertEqual(self.client.get(f"/api/analyses/{first_id}/result").json(), rejected.json())
        self.assertEqual(self.client.patch(
            f"/api/analyses/{first_id}/findings/absent", json={"status": "accepted"},
        ).status_code, 404)

        report = self.client.get(f"/api/analyses/{first_id}/report.docx")
        document = Document(BytesIO(report.content))
        self.assertIn("Служебная записка", [p.text for p in document.paragraphs])
        self.assertTrue(any("После: demo_after.docx" in p.text for p in document.paragraphs))
        self.assertTrue(any("Принято" in cell.text for row in document.tables[0].rows for cell in row.cells))

    def test_upload_validation_and_pending_result(self) -> None:
        response = self.client.post(
            "/api/analyses",
            files=[
                ("before", ("before.docx", docx_bytes("3.4. Старый пункт"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
                ("after", ("after.docx", docx_bytes("3.4. Новый пункт"), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ],
        )
        self.assertEqual(response.status_code, 200)
        analysis_id = response.json()["id"]
        self.assertEqual(self.client.get(f"/api/analyses/{analysis_id}").json()["status"], "done")

        bad = self.client.post(
            "/api/analyses",
            files=[("before", ("bad.txt", b"bad")), ("after", ("ok.docx", b"ok"))],
        )
        self.assertEqual(bad.status_code, 415)
        self.assertIn("error", bad.json())
        oversized = self.client.post("/api/analyses", files=[
            ("before", ("large.pdf", b"x" * (20 * 1024 * 1024 + 1))),
            ("after", ("after.docx", docx_bytes("3.4. Новый пункт"))),
        ])
        self.assertEqual(oversized.status_code, 413)
        self.assertEqual(self.client.post("/api/analyses", files=[
            ("before", ("before.docx", docx_bytes("3.4. Старый пункт"))),
        ]).status_code, 400)

        pending_id = self.client.app.state.store.create([
            ("before", "before.docx", docx_bytes("3.4. Старый пункт")),
            ("after", "after.docx", docx_bytes("3.4. Новый пункт")),
        ])
        pending = self.client.get(f"/api/analyses/{pending_id}/result")
        self.assertEqual(pending.status_code, 409)
        self.assertIn("error", pending.json())
        self.assertEqual(self.client.get(f"/api/analyses/{pending_id}/report.docx").status_code, 409)
        self.assertEqual(self.client.patch(
            f"/api/analyses/{pending_id}/findings/f1", json={"status": "accepted"},
        ).status_code, 409)
        missing = self.client.get("/api/analyses/" + "0" * 32)
        self.assertEqual(missing.status_code, 404)

    def test_analysis_failure_is_exposed_as_status(self) -> None:
        config = Settings(_env_file=None, ai_mock=False, storage_dir=Path(self.temp.name))
        client = TestClient(create_app(config))
        with patch("app.services.analyses.run_analysis", side_effect=RuntimeError("secret diagnostics")):
            response = client.post(
                "/api/analyses",
                files=[
                    ("before", ("before.docx", docx_bytes("3.4. До"))),
                    ("after", ("after.docx", docx_bytes("3.4. После"))),
                ],
            )
        self.assertEqual(response.status_code, 200)
        status = client.get(f"/api/analyses/{response.json()['id']}").json()
        self.assertEqual(status["status"], "failed")
        self.assertNotIn("secret diagnostics", str(status))
        client.close()

    def test_methods_cors_and_error_shapes(self) -> None:
        analysis_id = self.client.post("/api/analyses/demo").json()["id"]
        base = f"/api/analyses/{analysis_id}"
        wrong_method = self.client.post(base)
        self.assertEqual(wrong_method.status_code, 405)
        self.assertEqual(set(wrong_method.json()), {"error"})
        self.assertIn("GET", wrong_method.headers["allow"])
        for response in (
            self.client.get("/missing-route"),
            self.client.get("/api/analyses/" + "0" * 32),
            self.client.get(f"{base}/documents/missing/clauses/3.4"),
            self.client.get(f"{base}/documents/after-1/clauses/99.9"),
            self.client.patch(f"{base}/findings/missing", json={"status": "accepted"}),
        ):
            self.assertEqual(response.status_code, 404)
            self.assertEqual(set(response.json()), {"error"})
        invalid = self.client.patch(f"{base}/findings/f1", json={"status": "maybe"})
        self.assertEqual(invalid.status_code, 422)
        self.assertEqual(set(invalid.json()), {"error"})
        self.assertEqual(self.client.get(f"{base}/result").status_code, 200)

        allowed = self.client.options(
            "/api/analyses", headers={
                "Origin": "http://localhost:5173", "Access-Control-Request-Method": "POST",
            },
        )
        self.assertEqual(allowed.status_code, 200)
        self.assertEqual(allowed.headers["access-control-allow-origin"], "http://localhost:5173")
        blocked = self.client.options(
            "/api/analyses", headers={
                "Origin": "https://other.example", "Access-Control-Request-Method": "POST",
            },
        )
        self.assertNotIn("access-control-allow-origin", blocked.headers)

    def test_upload_multiple_formats_filenames_and_empty_files(self) -> None:
        response = self.client.post("/api/analyses", files=[
            ("before", ("C:\\docs\\old.DOCX", docx_bytes("3.4. До"))),
            ("before", ("other.xlsx", b"sheet")),
            ("after", ("../new.pdf", b"%PDF")),
        ])
        self.assertEqual(response.status_code, 200)
        folder = self.client.app.state.store.directory(response.json()["id"])
        self.assertEqual(
            [(d["name"], d["side"]) for d in self.client.app.state.store.documents(folder)],
            [("old.DOCX", "before"), ("other.xlsx", "before"), ("new.pdf", "after")],
        )
        self.assertTrue((folder / "before-1" / "old.DOCX").exists())
        self.assertTrue((folder / "before-2" / "other.xlsx").exists())
        self.assertTrue((folder / "after-1" / "new.pdf").exists())

        for files, code in (
            ([('before', ('empty.docx', b'')), ('after', ('ok.pdf', b'%PDF'))], 400),
            ([('before', ('old.doc', b'old')), ('after', ('ok.pdf', b'%PDF'))], 415),
            ([('after', ('ok.pdf', b'%PDF'))], 400),
        ):
            response = self.client.post("/api/analyses", files=files)
            self.assertEqual(response.status_code, code)
            self.assertEqual(set(response.json()), {"error"})

        accepted_limit = self.client.post("/api/analyses", files=[
            ("before", ("exact.pdf", b"x" * (20 * 1024 * 1024))),
            ("after", ("ok.pdf", b"%PDF")),
        ])
        self.assertEqual(accepted_limit.status_code, 200)

    def test_real_result_fields_parsed_source_and_file_fallback(self) -> None:
        config = Settings(_env_file=None, ai_mock=False, storage_dir=Path(self.temp.name))
        client = TestClient(create_app(config))
        result = json.loads((BACKEND_DIR / "mocks" / "result.json").read_text(encoding="utf-8"))
        result["meta"]["provider"] = "test"
        result["meta"]["documents"][0] = {"doc_id": "b1", "name": "before.docx", "side": "before"}
        result["documents"][0] = {
            "doc_id": "b1", "name": "before.docx", "side": "before",
            "clauses": [{"clause_id": "2.4.7", "section": "2", "text": "2.4.7. Текст из результата ИИ"}],
        }
        result["future_field"] = {"still_here": True}
        with patch("app.services.analyses.run_analysis", return_value=result):
            response = client.post("/api/analyses", files=[
                ("before", ("before.docx", docx_bytes("3.4. Источник до"))),
                ("after", ("after.docx", docx_bytes("3.4. Источник после"))),
            ])
        analysis_id = response.json()["id"]
        base = f"/api/analyses/{analysis_id}"
        self.assertEqual(client.get(f"{base}/result").json()["future_field"], {"still_here": True})
        self.assertEqual(
            client.get(f"{base}/documents/b1/clauses/2.4.7").json()["text"],
            "2.4.7. Текст из результата ИИ",
        )

        with patch("app.services.analyses.rebuild_conclusion", return_value="Пересобрано") as rebuild:
            reviewed = client.patch(f"{base}/findings/f1", json={"status": "accepted"})
        self.assertEqual(reviewed.status_code, 200)
        self.assertEqual(reviewed.json()["conclusion_md"], "Пересобрано")
        self.assertFalse(rebuild.call_args.kwargs["mock"])
        self.assertEqual(client.get(f"{base}/result").json()["future_field"], {"still_here": True})

        folder = client.app.state.store.directory(analysis_id)
        stored = deepcopy(client.app.state.store.result(folder))
        stored["documents"] = []
        client.app.state.store._write_json(folder / "result.json", stored)
        self.assertIn("Источник до", client.get(f"{base}/documents/b1/clauses/3.4").json()["text"])
        self.assertEqual(client.get(f"{base}/documents/b1/clauses/9.9").status_code, 404)
        metadata = client.app.state.store.documents(folder)
        metadata[0]["path"] = "absent.docx"
        client.app.state.store._write_json(folder / "documents.json", metadata)
        self.assertEqual(client.get(f"{base}/documents/b1/clauses/3.4").status_code, 422)
        client.close()

    def test_unexpected_error_is_not_exposed(self) -> None:
        client = TestClient(self.client.app, raise_server_exceptions=False)
        with patch.object(self.client.app.state.store, "history", side_effect=RuntimeError("secret diagnostics")):
            with patch("app.main.logging.exception"):
                response = client.get("/api/analyses")
        self.assertEqual(response.status_code, 500)
        self.assertEqual(response.json(), {"error": "Внутренняя ошибка сервера"})
        client.close()
