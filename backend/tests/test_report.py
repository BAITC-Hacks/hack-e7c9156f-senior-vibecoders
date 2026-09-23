from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase

from docx import Document

from app.services.analyses import build_report


class ReportTest(TestCase):
    def test_docx_contains_markdown_lists_evidence_reviews_and_disclaimer(self) -> None:
        with TemporaryDirectory() as temp:
            path = Path(temp) / "report.docx"
            result = {
                "conclusion_md": "# Итог\n\nАбзац заключения.\n- Первый пункт\n1. Второй пункт",
                "findings": [{
                    "type": "structure_change", "severity": "high", "title": "Изменение",
                    "evidence": [{"doc_name": "before.docx", "clause_id": "3.4", "quote": "Точная цитата"}],
                    "review": {"status": "rejected", "comment": "Не подтверждено"},
                }],
                "rejected_findings": [{
                    "type": "function_loss", "severity": "low", "title": "Спорный вывод",
                    "evidence": [], "review": {"status": "accepted", "comment": "Проверено"},
                }],
            }
            build_report(
                result,
                [{"name": "before.docx", "side": "before"}, {"name": "after.pdf", "side": "after"}],
                "2026-09-23T08:00:00+00:00", path,
            )

            document = Document(path)
            paragraphs = [p.text for p in document.paragraphs]
            self.assertIn("Служебная записка", paragraphs)
            self.assertIn("Дата: 23.09.2026", paragraphs)
            self.assertIn("До: before.docx", paragraphs)
            self.assertIn("После: after.pdf", paragraphs)
            self.assertIn("Абзац заключения.", paragraphs)
            self.assertIn("Первый пункт", paragraphs)
            self.assertIn("Второй пункт", paragraphs)
            self.assertIn("Выводы носят рекомендательный характер и требуют проверки ответственным сотрудником.", paragraphs)
            self.assertEqual(next(p.style.name for p in document.paragraphs if p.text == "Первый пункт"), "List Bullet")
            self.assertEqual(next(p.style.name for p in document.paragraphs if p.text == "Второй пункт"), "List Number")

            self.assertEqual(len(document.tables), 1)
            rows = [[cell.text for cell in row.cells] for row in document.tables[0].rows]
            self.assertEqual(rows[0], ["Тип", "Критичность", "Заголовок", "Источники", "Проверка"])
            self.assertIn("before.docx, п. 3.4 — «Точная цитата»", rows[1][3])
            self.assertIn("Отклонено: Не подтверждено", rows[1][4])
            self.assertIn("Принято: Проверено; опровергнуто критиком", rows[2][4])
            self.assertFalse(list(Path(temp).glob("*.tmp")))
