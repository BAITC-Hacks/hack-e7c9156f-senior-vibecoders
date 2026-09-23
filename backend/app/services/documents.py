import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


CLAUSE_START = re.compile(r"^\s*(\d+(?:\.\d+)*(?:\.[A-Za-zА-Яа-яЁё])?)(?:\.)?\s+(.+)")


def document_lines(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = Document(path)
        lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(cell.text.strip() for cell in row.cells))
        return lines
    if suffix == ".pdf":
        return [line.strip() for page in PdfReader(path).pages for line in (page.extract_text() or "").splitlines() if line.strip()]
    if suffix == ".xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            return [" ".join(str(cell) for cell in row if cell is not None).strip()
                    for sheet in book.worksheets for row in sheet.iter_rows(values_only=True)
                    if any(cell is not None for cell in row)]
        finally:
            book.close()
    raise ValueError("Неподдерживаемый формат документа")


def get_clause(path: Path, clause_id: str) -> str | None:
    current_id = None
    collected: list[str] = []
    for line in document_lines(path):
        match = CLAUSE_START.match(line)
        if match:
            if current_id == clause_id:
                break
            current_id = match.group(1)
        if current_id == clause_id:
            collected.append(line)
    return "\n".join(collected) or None
