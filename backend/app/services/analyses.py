import json
import os
from pathlib import Path
from uuid import uuid4

from docx import Document
from app.clients.ai import run_analysis
from app.config import Settings
from app.schemas import AnalysisResult, AnalysisStatus


DEMO_BEFORE = [
    "3.4. В БВА входят ДНМ, ДККМ и Директор направления внутреннего аудита.",
    "3.8. В ДККМ предусмотрены Директор по контролю качества аудита и методологии и Менеджер по аудиту.",
]
DEMO_AFTER = [
    "3.4. В БВА входят ДНМ, ДККМ, ДИТААД и ДОА.",
    "3.9. В ДККМ предусмотрены обновлённые должности.",
]


class AnalysisStore:
    def __init__(self, settings: Settings):
        self.settings = settings
        self.root = settings.storage_dir.resolve()
        self.root.mkdir(parents=True, exist_ok=True)

    def directory(self, analysis_id: str) -> Path | None:
        if len(analysis_id) != 32 or any(c not in "0123456789abcdef" for c in analysis_id):
            return None
        path = self.root / analysis_id
        return path if path.is_dir() else None

    @staticmethod
    def _write_json(path: Path, value: object) -> None:
        temporary = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
        temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
        os.replace(temporary, path)

    @staticmethod
    def _read_json(path: Path) -> object:
        return json.loads(path.read_text(encoding="utf-8"))

    def create(self, files: list[tuple[str, str, bytes]], *, demo: bool = False) -> str:
        analysis_id = uuid4().hex
        folder = self.root / analysis_id
        folder.mkdir()
        documents = []
        counts = {"before": 0, "after": 0}
        for side, name, content in files:
            counts[side] += 1
            doc_id = f"{side}-{counts[side]}"
            path = folder / f"{doc_id}{Path(name).suffix.lower()}"
            path.write_bytes(content)
            documents.append({"doc_id": doc_id, "name": name, "side": side, "path": path.name})
        self._write_json(folder / "documents.json", documents)
        self._write_json(folder / "status.json", AnalysisStatus(id=analysis_id, status="queued", progress=0).model_dump(exclude_none=True))
        if demo:
            (folder / "demo").touch()
        return analysis_id

    def create_demo(self) -> str:
        files = []
        for side, name, lines in [
            ("before", "demo_before.docx", DEMO_BEFORE),
            ("after", "demo_after.docx", DEMO_AFTER),
        ]:
            from io import BytesIO

            doc = Document()
            for line in lines:
                doc.add_paragraph(line)
            output = BytesIO()
            doc.save(output)
            files.append((side, name, output.getvalue()))
        return self.create(files, demo=True)

    def status(self, folder: Path) -> AnalysisStatus:
        return AnalysisStatus.model_validate(self._read_json(folder / "status.json"))

    def documents(self, folder: Path) -> list[dict]:
        return self._read_json(folder / "documents.json")

    def result(self, folder: Path) -> AnalysisResult:
        return AnalysisResult.model_validate(self._read_json(folder / "result.json"))

    def run(self, analysis_id: str) -> None:
        folder = self.directory(analysis_id)
        if folder is None:
            return
        previous = self.status(folder)
        self._write_json(folder / "status.json", previous.model_copy(update={"status": "running", "step": "parsing", "progress": 0.01}).model_dump(exclude_none=True))

        def progress(step: str, value: float) -> None:
            status = AnalysisStatus(id=analysis_id, status="running", step=step, progress=value)
            self._write_json(folder / "status.json", status.model_dump(exclude_none=True))

        documents = self.documents(folder)
        before = [folder / d["path"] for d in documents if d["side"] == "before"]
        after = [folder / d["path"] for d in documents if d["side"] == "after"]
        try:
            result = run_analysis(before, after, progress, mock=self.settings.ai_mock)
        except Exception as exc:
            if (folder / "demo").exists() and not self.settings.ai_mock:
                try:
                    result = run_analysis(before, after, progress, mock=True)
                except Exception as fallback_exc:
                    self._fail(folder, analysis_id, fallback_exc)
                    return
            else:
                self._fail(folder, analysis_id, exc)
                return

        self._write_json(folder / "result.json", result.model_dump(exclude_none=True))
        self._write_json(folder / "status.json", AnalysisStatus(id=analysis_id, status="done", step="report", progress=1).model_dump(exclude_none=True))

    def _fail(self, folder: Path, analysis_id: str, exc: Exception) -> None:
        message = "Пакет ИИ недоступен" if isinstance(exc, ModuleNotFoundError) else "Не удалось выполнить анализ"
        self._write_json(folder / "status.json", AnalysisStatus(id=analysis_id, status="failed", progress=0, error=message).model_dump(exclude_none=True))


def build_report(result: AnalysisResult, path: Path) -> None:
    doc = Document()
    doc.add_heading("Заключение по анализу", 0)
    for line in result.conclusion_md.splitlines():
        if line.strip():
            doc.add_paragraph(line.lstrip("# "))
    if result.findings:
        doc.add_heading("Выводы", level=1)
        for finding in result.findings:
            doc.add_heading(finding.title, level=2)
            doc.add_paragraph(finding.description)
            if finding.recommendation:
                doc.add_paragraph(f"Рекомендация: {finding.recommendation}")
            for evidence in finding.evidence:
                doc.add_paragraph(f"Источник: {evidence.doc_name}, п. {evidence.clause_id}: {evidence.quote}")
    doc.save(path)
