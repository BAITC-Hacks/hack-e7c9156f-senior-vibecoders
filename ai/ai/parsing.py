"""Чтение DOCX/PDF/XLSX и нарезка текста на пункты с clause_id. Без LLM.

clause_id:
  "3"      — заголовок раздела
  "3.4"    — нумерованный пункт
  "3.4.а"  — буквенный подпункт пункта 3.4
  "Лист1:12" — строка таблицы Excel
  "т1.3"   — строка 3 таблицы 1 в Word (напр. приложение с оргструктурой)
  "абз.5"  — абзац 5 документа без нумерации (страховка)
  "0"      — текст до первого пункта (шапка документа)

Автонумерация Word («1.1.» и «а)», которых нет в тексте абзаца) восстанавливается по numbering.xml.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from .schemas import Clause, DocumentText, Side

# "3.4. Текст", "3.10.Текст", "1. Общие положения"
_NUM_RE = re.compile(r"^(\d+(?:\.\d+)*)\.?\s*(.*)$", re.S)
# "а. Текст", "б) Текст"
_LETTER_RE = re.compile(r"^([а-яё])[.)]\s+(.*)$", re.S | re.I)
# пункт, склеенный внутри абзаца: "...направления. 3.10.Рабочие места..."
_INLINE_NUM_RE = re.compile(r"(?<=[.;:])\s+(?=\d+\.\d+(?:\.\d+)*\.\s*[А-ЯЁA-Z])")
# строка оглавления: "3. СТРУКТУРА И ОРГАНИЗАЦИЯ РАБОТЫ ВНУТРЕННЕГО АУДИТА 8"
_TOC_RE = re.compile(r"^\d+\.\s+[А-ЯЁA-Z0-9\s,.«»\"()\-–]+\s\d+\s*$")
_SKIP_LINES = {"оглавление", "содержание", "приложения"}
# "Приложение 1. Кодекс этики …"
_APPENDIX_RE = re.compile(r"^Приложение\s+(\d+)\b", re.I)


@dataclass
class TableRow:
    """Строка таблицы — отдельный пункт с id «т<таблица>.<строка>»."""

    table: int
    row: int
    text: str


Line = str | TableRow

_RU_LETTERS = "абвгдежзиклмнопрстуфхцчшщэюя"  # как в Word (без ё, й, ъ, ы, ь)


def _fmt_number(n: int, fmt: str) -> str:
    if fmt in ("lowerLetter", "upperLetter"):
        s = chr(ord("a") + (n - 1) % 26)
        return s.upper() if fmt == "upperLetter" else s
    if fmt in ("russianLower", "russianUpper"):
        s = _RU_LETTERS[(n - 1) % len(_RU_LETTERS)]
        return s.upper() if fmt == "russianUpper" else s
    if fmt in ("lowerRoman", "upperRoman"):
        vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
        s, x = "", n
        for v, r in vals:
            while x >= v:
                s, x = s + r, x - v
        return s.upper() if fmt == "upperRoman" else s
    return str(n)  # decimal и всё остальное


class _Numbering:
    """Восстанавливает автонумерацию Word («1.», «1.1.», «а)»), которой нет в тексте абзаца."""

    def __init__(self, doc):
        from docx.oxml.ns import qn

        self.qn = qn
        self.levels: dict[str, dict[int, tuple[int, str, str]]] = {}  # abstractId → ilvl → (start, fmt, text)
        self.num_to_abs: dict[str, str] = {}
        self.counters: dict[str, dict[int, int]] = {}
        try:
            root = doc.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return
        for an in root.findall(qn("w:abstractNum")):
            lv = {}
            for l in an.findall(qn("w:lvl")):
                def val(tag, default):
                    e = l.find(qn(tag))
                    return e.get(qn("w:val")) if e is not None else default
                lv[int(l.get(qn("w:ilvl")))] = (int(val("w:start", "1")), val("w:numFmt", "decimal"), val("w:lvlText", ""))
            self.levels[an.get(qn("w:abstractNumId"))] = lv
        for n in root.findall(qn("w:num")):
            a = n.find(qn("w:abstractNumId"))
            if a is not None:
                self.num_to_abs[n.get(qn("w:numId"))] = a.get(qn("w:val"))

    def _num_pr(self, p):
        qn = self.qn
        ppr = p._p.pPr
        if ppr is not None and ppr.numPr is not None and ppr.numPr.numId is not None:
            ilvl = ppr.numPr.ilvl.val if ppr.numPr.ilvl is not None else 0
            return str(ppr.numPr.numId.val), int(ilvl)
        style = p.style
        while style is not None:  # нумерация через стиль (в т.ч. унаследованный)
            spr = style.element.find(qn("w:pPr"))
            npr = spr.find(qn("w:numPr")) if spr is not None else None
            if npr is not None and npr.find(qn("w:numId")) is not None:
                il = npr.find(qn("w:ilvl"))
                return npr.find(qn("w:numId")).get(qn("w:val")), int(il.get(qn("w:val"))) if il is not None else 0
            style = style.base_style
        return None

    def label(self, p) -> str:
        pr = self._num_pr(p)
        if not pr or pr[0] == "0" or pr[0] not in self.num_to_abs:
            return ""
        abs_id = self.num_to_abs[pr[0]]
        levels = self.levels.get(abs_id, {})
        ilvl = pr[1]
        if ilvl not in levels or levels[ilvl][1] in ("bullet", "none"):
            return ""
        cnt = self.counters.setdefault(abs_id, {})
        cnt[ilvl] = cnt.get(ilvl, levels[ilvl][0] - 1) + 1
        for deeper in [k for k in cnt if k > ilvl]:
            del cnt[deeper]
        text = levels[ilvl][2]
        for k in range(ilvl + 1):
            start, fmt, _ = levels.get(k, (1, "decimal", ""))
            text = text.replace(f"%{k + 1}", _fmt_number(cnt.get(k, start), fmt))
        return text.strip()


def _docx_lines(path: Path) -> list[Line]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    numbering = _Numbering(d)
    lines: list[Line] = []
    n_table = 0
    for el in d.element.body.iterchildren():
        tag = el.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(el, d)
            text, label = p.text, numbering.label(p)
            if label and text.strip() and not text.lstrip().startswith(label):
                text = f"{label} {text.lstrip()}"
            lines.append(text)
        elif tag == "tbl":
            n_table += 1
            for r, row in enumerate(Table(el, d).rows, 1):
                cells: list[str] = []
                for c in row.cells:  # объединённые ячейки python-docx возвращает повторно
                    t = c.text.strip()
                    if t and (not cells or cells[-1] != t):
                        cells.append(t)
                if cells:
                    lines.append(TableRow(n_table, r, " | ".join(cells)))
    return lines


def _pdf_lines(path: Path) -> list[str]:
    import fitz  # pymupdf

    lines: list[str] = []
    with fitz.open(str(path)) as doc:
        for page in doc:
            lines.extend(page.get_text().splitlines())
    return lines


def _xlsx_clauses(path: Path) -> list[Clause]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    clauses: list[Clause] = []
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if cells:
                clauses.append(Clause(clause_id=f"{ws.title}:{i}", section=ws.title, text=" | ".join(cells)))
    return clauses


def _clean(text: str) -> str:
    text = text.replace("\xa0", " ").replace("\t", " ")
    return re.sub(r"[ ]{2,}", " ", text).strip()


def _section(cid: str) -> str:
    if cid.startswith("прил."):
        return "прил"
    return cid.split(".", 1)[0].split("#", 1)[0]  # «3.4.а» → «3», «т1.2» → «т1», «абз.5» → «абз»


def segment(lines: list[Line]) -> list[Clause]:
    """Нарезка строк на пункты по нумерации (текстовой или восстановленной из автонумерации Word)."""
    clauses: list[Clause] = []
    seen: dict[str, int] = {}
    current_num: str | None = None  # последний нумерованный пункт — родитель для «а.», «б.»

    def add(cid: str, text: str) -> None:
        if cid in seen:  # повтор номера (ошибка в документе) — не теряем текст
            seen[cid] += 1
            cid = f"{cid}#{seen[cid]}"
        else:
            seen[cid] = 1
        clauses.append(Clause(clause_id=cid, section=_section(cid), text=text))

    for raw in lines:
        if isinstance(raw, TableRow):
            add(f"т{raw.table}.{raw.row}", _clean(raw.text))
            current_num = None
            continue
        raw = _clean(raw)
        if not raw or _TOC_RE.match(raw) or raw.lower().rstrip(".:") in _SKIP_LINES:
            continue
        m = _APPENDIX_RE.match(raw)
        if m:
            current_num = None
            add(f"прил.{m.group(1)}", raw)
            continue
        for part in _INLINE_NUM_RE.split(raw):
            part = part.strip()
            if not part:
                continue
            m = _NUM_RE.match(part)
            if m and m.group(2):
                current_num = m.group(1)
                add(current_num, part)
                continue
            m = _LETTER_RE.match(part)
            if m and current_num:
                add(f"{current_num}.{m.group(1).lower()}", part)
                continue
            if clauses:  # продолжение предыдущего пункта (маркированный список, перенос строки)
                clauses[-1] = clauses[-1].model_copy(update={"text": f"{clauses[-1].text}\n{part}"})
            else:
                add("0", part)
    return clauses


def _numbered_count(clauses: list[Clause]) -> int:
    return sum(1 for c in clauses if c.clause_id[0].isdigit())


def paragraph_clauses(lines: list[Line]) -> list[Clause]:
    """Страховка для документов без нумерации: пункт = абзац («абз.5»), таблицы — «т1.2»."""
    out: list[Clause] = []
    n = 0
    for raw in lines:
        if isinstance(raw, TableRow):
            cid = f"т{raw.table}.{raw.row}"
            out.append(Clause(clause_id=cid, section=_section(cid), text=_clean(raw.text)))
            continue
        text = _clean(raw)
        if text:
            n += 1
            out.append(Clause(clause_id=f"абз.{n}", section="абз", text=text))
    return out


def _segment_or_paragraphs(lines: list[Line]) -> list[Clause]:
    clauses = segment(lines)
    text_lines = sum(1 for l in lines if isinstance(l, str) and _clean(l))
    if _numbered_count(clauses) < 2 and text_lines >= 3:  # нумерации нет — ссылаемся на абзацы
        return paragraph_clauses(lines)
    return clauses


def parse_document(path: str | Path, side: Side, doc_id: str) -> DocumentText:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".docx":
        clauses = _segment_or_paragraphs(_docx_lines(path))
    elif ext == ".pdf":
        clauses = _segment_or_paragraphs(_pdf_lines(path))
    elif ext in (".xlsx", ".xlsm"):
        clauses = _xlsx_clauses(path)
    else:
        raise ValueError(f"Неподдерживаемый формат: {path.name} (нужен .docx, .pdf или .xlsx)")
    return DocumentText(doc_id=doc_id, name=path.name, side=side, clauses=clauses)


def parse_documents(before: list[Path], after: list[Path]) -> list[DocumentText]:
    docs = [parse_document(p, "before", f"b{i}") for i, p in enumerate(before, 1)]
    docs += [parse_document(p, "after", f"a{i}") for i, p in enumerate(after, 1)]
    return docs


def clause_text(doc: DocumentText, clause_id: str, with_children: bool = True) -> str | None:
    """Текст пункта; с подпунктами (3.4 → 3.4 + 3.4.а + 3.4.б …)."""
    parts = [
        c.text
        for c in doc.clauses
        if c.clause_id == clause_id or (with_children and c.clause_id.startswith(clause_id + "."))
    ]
    return "\n".join(parts) if parts else None


def render_for_prompt(doc: DocumentText, sections: set[str] | None = None) -> str:
    """Документ в виде «[clause_id] текст» — формат, в котором LLM видит и цитирует пункты."""
    rows = [f"[{c.clause_id}] {c.text}" for c in doc.clauses if sections is None or c.section in sections]
    return f"=== Документ {doc.doc_id}: {doc.name} ({'до' if doc.side == 'before' else 'после'}) ===\n" + "\n".join(rows)
