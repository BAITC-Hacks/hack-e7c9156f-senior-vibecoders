from pathlib import Path

import pytest

from ai.parsing import clause_text, parse_document, segment

SAMPLES = Path(__file__).resolve().parents[2] / "samples"


def ids(clauses):
    return [c.clause_id for c in clauses]


def test_numbered_and_letter_items():
    cl = segment([
        "1. Общие положения",
        "3.4. БВА состоит из следующих структурных подразделений:",
        "а. Департамент ИТ-аудита (ДИТААД).",
        "б. Департамент операционного аудита (ДОА).",
        "- продолжение списка",
    ])
    assert ids(cl) == ["1", "3.4", "3.4.а", "3.4.б"]
    assert cl[1].section == "3"
    assert cl[-1].text.endswith("- продолжение списка")


def test_inline_clause_split():
    cl = segment(["3.9. Состав: б. Руководитель направления. 3.10.Рабочие места могут располагаться в филиалах."])
    assert ids(cl) == ["3.9", "3.10"]
    assert cl[1].text.startswith("3.10.Рабочие места")


def test_toc_and_appendix():
    cl = segment(["13.3. Последний пункт.", "Оглавление", "1. ОБЩИЕ ПОЛОЖЕНИЯ 1", "Приложение 1. Кодекс этики"])
    assert ids(cl) == ["13.3", "прил.1"]
    assert cl[0].text == "13.3. Последний пункт."


def test_duplicate_ids_kept():
    cl = segment(["9.3. Пункт", "а. один", "вводный текст", "а. другой"])
    assert ids(cl) == ["9.3", "9.3.а", "9.3.а#2"]


def test_clause_text_with_children():
    cl = segment(["3.4. Состав:", "а. ДИТААД", "3.40. Другой пункт"])
    from ai.schemas import DocumentText

    doc = DocumentText(doc_id="a1", name="x", side="after", clauses=cl)
    assert clause_text(doc, "3.4") == "3.4. Состав:\nа. ДИТААД"  # 3.40 не считается подпунктом 3.4


def _docx_with_autonumbering(path):
    import docx
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    d = docx.Document()
    num = d.part.numbering_part.element
    num.insert(0, parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="90">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2."/></w:lvl>'
        '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="russianLower"/><w:lvlText w:val="%3)"/></w:lvl>'
        '</w:abstractNum>'))
    num.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="90"><w:abstractNumId w:val="90"/></w:num>'))
    d.add_paragraph("ПОЛОЖЕНИЕ об Отделе")
    for level, text in [(0, "Общие положения"), (1, "Отдел входит в ДИТ."), (0, "Функции"),
                        (1, "Обеспечивает работу серверов."), (1, "Управляет доступом:"),
                        (2, "выдаёт доступ;"), (2, "блокирует доступ."), (0, "Ответственность")]:
        p = d.add_paragraph(text)
        p._p.get_or_add_pPr().append(parse_xml(
            f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{level}"/><w:numId w:val="90"/></w:numPr>'))
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Подразделение", "Сокращение"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Отдел инфраструктуры", "ОИ"
    d.save(str(path))


def test_word_autonumbering_and_tables(tmp_path):
    p = tmp_path / "auto.docx"
    _docx_with_autonumbering(p)
    doc = parse_document(p, "after", "a1")
    assert ids(doc.clauses) == ["0", "1", "1.1", "2", "2.1", "2.2", "2.2.а", "2.2.б", "3", "т1.1", "т1.2"]
    assert doc.clauses[5].text.startswith("2.2. Управляет доступом")
    assert "ОИ" in clause_text(doc, "т1.2")


def test_fallback_to_paragraph_ids_without_numbering():
    from ai.parsing import TableRow, _segment_or_paragraphs

    cl = _segment_or_paragraphs(["ПРИКАЗ", "Об утверждении структуры", "Утвердить структуру согласно приложению.",
                                 TableRow(1, 1, "Отдел | ОИ")])
    assert ids(cl) == ["абз.1", "абз.2", "абз.3", "т1.1"]


@pytest.mark.skipif(not (SAMPLES / "after").exists(), reason="нет samples/")
def test_real_document_r9():
    doc = parse_document(next((SAMPLES / "after").glob("*.docx")), "after", "a1")
    t = clause_text(doc, "3.4")
    assert "ДИТААД" in t and "ДОА" in t and "ДНМ" in t and "ДККМ" in t
    assert clause_text(doc, "3.10", with_children=False).startswith("3.10.Рабочие места")
    assert "3.11" in ids(doc.clauses)
